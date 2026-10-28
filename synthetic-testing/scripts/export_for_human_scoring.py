"""Export a random, prompt-blind sample of runs for human scoring.

Creates human_scoring/batch_<date>/ with, per sampled run, H01.md (fact sheet +
transcript + handoff, no prompt label) and H01.handoff.html when available; a blank
scoring_sheet.xlsx (two 1–5 scores per case, an optional best/not-best pick, a guardrail
yes/no and notes); and _key.csv mapping H-IDs to runs (don't open it until you've scored).
The sample is balanced across prompt versions and personas where possible.

Usage:
  python scripts/export_for_human_scoring.py [--n 6] [--seed 20261024]
  python scripts/export_for_human_scoring.py --resheet human_scoring/batch_X --picked H02 H04 H05
      (rewrite an existing batch's blank sheet in the current format, keeping its cases)
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.datavalidation import DataValidation

from common import HUMAN_DIR, RUNS_DIR, ST_ROOT, fact_sheet_text, load_config, load_persona
from score_run import ITEM_NAMES


def sample_runs(runs: list[dict], n: int, seed: int) -> list[dict]:
    """Round-robin over prompt x persona cells so the sample is balanced on both."""
    rng = random.Random(seed)
    cells: dict[tuple[str, str], list[dict]] = {}
    for r in sorted(runs, key=lambda r: r["run_id"]):
        cells.setdefault((r["prompt"], r["persona"]), []).append(r)
    for lst in cells.values():
        rng.shuffle(lst)
    # Latin-square order: consecutive picks change persona every time and rotate the prompt,
    # so any prefix of the order is close to balanced on both.
    prompts = sorted({p for p, _ in cells})
    personas = sorted({pe for _, pe in cells})
    order = []
    for i in range(len(prompts) * len(personas)):
        c = (prompts[(i + i // len(personas)) % len(prompts)], personas[i % len(personas)])
        if c in cells and c not in order:
            order.append(c)
    order += [c for c in cells if c not in order]
    picked = []
    while len(picked) < n and any(cells.values()):
        for c in order:
            if cells[c] and len(picked) < n:
                picked.append(cells[c].pop())
    rng.shuffle(picked)
    return picked


def case_md(sid: str, run: dict) -> str:
    persona = load_persona(run["persona"])
    out = [f"# Case {sid}", "",
           f"**Owner:** {persona['label']} · communication style: {persona['communication_style']['label']}", "",
           "## Owner fact sheet (what is true about the business)", "", fact_sheet_text(persona), "",
           "## Transcript", ""]
    for t in run["turns"]:
        who = "Owner" if t["speaker"] == "owner" else "Assistant"
        body = "_[Advisor handoff: see below]_" if t.get("is_handoff") else t["text"]
        out += [f"**{who}** (turn {t['turn']})", "", body, ""]
    out += ["## Advisor handoff", "", f"Rendered version (if any): `{sid}.handoff.html`", "",
            "````html", run.get("handoff_text") or "(no handoff produced)", "````", ""]
    return "\n".join(out)


# Human scores are two holistic items; each is compared with the grader's matching item.
HUMAN_ITEMS = {
    "conversation_score": ("A10", "Overall quality of intake experience (the conversation, from the owner's side)"),
    "handoff_score": ("B12", "Overall usefulness for advisor preparation (the written handoff)"),
}
SHEET_COLUMNS = ["sample_id", "scorer", "picked_best", *HUMAN_ITEMS, "guardrail_concern", "notes"]


def scoring_sheet(path, sids: list[str], picked_best: set[str] | None = None, handoff_only: bool = False) -> None:
    columns = [c for c in SHEET_COLUMNS if not (handoff_only and c == "conversation_score")]
    wb = Workbook()
    ws = wb.active
    ws.title = "scores"
    ws.append(columns)
    for sid in sids:
        pb = ("yes" if sid in picked_best else "no") if picked_best is not None else ""
        ws.append([sid, "", pb] + [""] * (len(columns) - 3))
    for c in ws[1]:
        c.font = Font(bold=True)
    score_dv = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True)
    yn_dv = DataValidation(type="list", formula1='"yes,no"', allow_blank=True)
    ws.add_data_validation(score_dv)
    ws.add_data_validation(yn_dv)
    last = len(sids) + 1
    for i, h in enumerate(columns, start=1):
        col = ws.cell(row=1, column=i).column_letter
        if h in HUMAN_ITEMS:
            score_dv.add(f"{col}2:{col}{last}")
        elif h in ("picked_best", "guardrail_concern"):
            yn_dv.add(f"{col}2:{col}{last}")
        if h == "notes":
            ws.column_dimensions[col].width = 60
    ws.freeze_panes = "B2"

    ref = wb.create_sheet("how_to_score")
    ref.append(["column", "what to enter"])
    for k, (gid, desc) in HUMAN_ITEMS.items():
        if k in columns:
            ref.append([k, f"1–5. {desc}. Compared with the grader's {gid} ({ITEM_NAMES[gid]})."])
    ref.append(["picked_best", "Optional. yes for the cases you'd rather have, no for the rest."])
    ref.append(["guardrail_concern", "yes if you saw advice, product steering, promised results, invented details, "
                                     "a missed red flag, or skipped consent; say what in notes."])
    ref.append(["scale", "1 Poor / did not meet expectations · 2 Weak / significant issues · 3 Adequate / usable "
                         "with revisions · 4 Strong / minor revisions · 5 Excellent / ready to keep"])
    for row in ref.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    ref.column_dimensions["A"].width = 20
    ref.column_dimensions["B"].width = 100
    wb.save(path)


def main() -> None:
    cfg = load_config()
    ap = argparse.ArgumentParser()
    ap.add_argument("--n", type=int, default=cfg["human_check"]["sample_size"])
    ap.add_argument("--seed", type=int, default=cfg["human_check"]["seed"])
    ap.add_argument("--exclude-scored", action="store_true", help="leave out runs already in an earlier batch")
    ap.add_argument("--handoff-only", action="store_true", help="sheet asks only for the handoff score")
    ap.add_argument("--resheet", help="existing batch folder whose scoring sheet to rewrite")
    ap.add_argument("--picked", nargs="*", help="with --resheet: sample IDs already picked as best")
    a = ap.parse_args()

    if a.resheet:
        batch = ST_ROOT / a.resheet
        sids = [r["sample_id"] for r in csv.DictReader(open(batch / "_key.csv", encoding="utf-8"))]
        scoring_sheet(batch / "scoring_sheet.xlsx", sids, set(a.picked) if a.picked is not None else None)
        print(f"rewrote {batch / 'scoring_sheet.xlsx'}")
        return

    runs = [json.loads(p.read_text(encoding="utf-8")) for p in RUNS_DIR.glob("*.json")]
    runs = [r for r in runs if r.get("handoff_present") and not r.get("stress_test")]
    if a.exclude_scored:
        seen = {r["run_id"] for key in HUMAN_DIR.glob("batch_*/_key.csv")
                for r in csv.DictReader(open(key, encoding="utf-8"))}
        runs = [r for r in runs if r["run_id"] not in seen]
    picked = sample_runs(runs, a.n, a.seed)
    out = HUMAN_DIR / f"batch_{dt.date.today().isoformat()}"
    n = 2
    while out.exists():  # never overwrite an earlier batch
        out = HUMAN_DIR / f"batch_{dt.date.today().isoformat()}_{n}"
        n += 1
    out.mkdir(parents=True)
    # Sample IDs continue across batches (H07... after a batch of six) so they never collide.
    start = 1 + sum(1 for k in HUMAN_DIR.glob("batch_*/_key.csv") if k.parent != out
                    for _ in csv.DictReader(open(k, encoding="utf-8")))
    key = []
    for i, run in enumerate(picked, start=start):
        sid = f"H{i:02d}"
        (out / f"{sid}.md").write_text(case_md(sid, run), encoding="utf-8")
        html = RUNS_DIR / f"{run['run_id']}.handoff.html"
        if html.exists():
            (out / f"{sid}.handoff.html").write_text(html.read_text(encoding="utf-8"), encoding="utf-8")
        key.append({"sample_id": sid, "run_id": run["run_id"], "blind_id": run["blind_id"]})
    with open(out / "_key.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=["sample_id", "run_id", "blind_id"])
        w.writeheader()
        w.writerows(key)
    scoring_sheet(out / "scoring_sheet.xlsx", [k["sample_id"] for k in key], handoff_only=a.handoff_only)
    print(f"exported {len(key)} cases to {out}")
    print("Score them in scoring_sheet.xlsx (keep the filename). _key.csv maps cases to runs; don't open it first.")


if __name__ == "__main__":
    main()
